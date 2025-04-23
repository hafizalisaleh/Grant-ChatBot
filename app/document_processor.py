import json
import re
from docxtpl import DocxTemplate
from docx import Document
import tempfile, zipfile
from fastapi import UploadFile, HTTPException
from .api_client import call_groq_chat

async def process_document(file: UploadFile, prompt: str) -> bytes:
    # 1) Save upload
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    # 2) Validate .docx
    if not zipfile.is_zipfile(tmp_path):
        raise HTTPException(400, "Not a valid .docx file.")
    with zipfile.ZipFile(tmp_path) as z:
        if 'word/document.xml' not in z.namelist():
            raise HTTPException(400, "Missing word/document.xml in .docx.")

    # 3) Extract raw text (for context)
    raw = Document(tmp_path)
    full_text = "\n".join(p.text for p in raw.paragraphs)

    # 4) Call the LLM once, asking it to return JSON for sections
    system = {
        "role": "system",
        "content": (
            "You are a precise editor and will output pure JSON without any commentary or repetition."
        )
    }
    user = {
        "role": "user",
        "content": (
            f"{prompt}\n\nOriginal Document Text:\n{full_text}\n\n"
            "Return a JSON object with keys matching the placeholders and values only the rewritten text."
        )
    }
    resp = call_groq_chat([system, user])
    content = resp["choices"][0]["message"]["content"].strip()

    # 5) Extract JSON block if needed
    if not content.startswith("{"):
        match = re.search(r"\{.*\}", content, flags=re.DOTALL)
        if match:
            content = match.group(0)

    # 6) Parse JSON or error
    try:
        new_sections = json.loads(content)
    except json.JSONDecodeError:
        raise HTTPException(500, f"Failed to parse JSON: {content}")

    # 7) Generic echo-removal helper
    def strip_repeated_suffix(text: str) -> str:
        tokens = text.split()
        for n in range(len(tokens)//2, 0, -1):
            if tokens[:n] == tokens[-n:]:
                return " ".join(tokens[:-n]).strip()
        return text

    # 8) Clean up each section
    for key, txt in new_sections.items():
        new_sections[key] = strip_repeated_suffix(txt)

    # 9) Render placeholders into template
    tpl = DocxTemplate(tmp_path)
    tpl.render(new_sections)

    # 10) Save and return
    out_path = tmp_path.replace(".docx", "-rendered.docx")
    tpl.save(out_path)
    with open(out_path, "rb") as f:
        return f.read()